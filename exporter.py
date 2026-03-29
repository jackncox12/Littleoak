"""Export prospects to Excel and CSV."""

import pandas as pd
from pathlib import Path
from datetime import datetime


COLUMN_ORDER = [
    "business_name",
    "address",
    "phone",
    "emails",
    "website",
    "contact_name",
    "linkedin",
    "estimated_size",
    "google_rating",
    "rating_count",
    "current_waste_provider_hints",
    "status",
    "google_maps_url",
    "business_types",
]

FRIENDLY_NAMES = {
    "business_name": "Business Name",
    "address": "Address",
    "phone": "Phone",
    "emails": "Emails",
    "website": "Website",
    "contact_name": "Contact Name(s)",
    "linkedin": "LinkedIn",
    "estimated_size": "Est. Size",
    "google_rating": "Google Rating",
    "rating_count": "No. of Reviews",
    "current_waste_provider_hints": "Current Waste Provider (Hints)",
    "status": "Status",
    "google_maps_url": "Google Maps",
    "business_types": "Business Types",
}


def export(prospects: list[dict], output_dir: str = ".", prefix: str = "prospects") -> dict[str, str]:
    """Export prospects list to both Excel and CSV. Returns paths."""
    if not prospects:
        raise ValueError("No prospects to export.")

    df = pd.DataFrame(prospects)

    # Ensure all expected columns exist
    for col in COLUMN_ORDER:
        if col not in df.columns:
            df[col] = ""

    df = df[COLUMN_ORDER]
    df = df.rename(columns=FRIENDLY_NAMES)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = Path(output_dir)
    base.mkdir(parents=True, exist_ok=True)

    excel_path = str(base / f"{prefix}_{timestamp}.xlsx")
    csv_path = str(base / f"{prefix}_{timestamp}.csv")

    # --- Excel with formatting ---
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Prospects")
        ws = writer.sheets["Prospects"]

        # Header styling
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Auto-fit columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        # Freeze header row
        ws.freeze_panes = "A2"

    # --- CSV ---
    df.to_csv(csv_path, index=False)

    return {"excel": excel_path, "csv": csv_path}
